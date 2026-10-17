"""
Shared constants and small helpers used across pages.
"""
import calendar
import datetime as dt

STATUS_OPTIONS = ["todo", "in_progress", "in_review", "done"]
STATUS_LABELS = {
    "todo": "To Do",
    "in_progress": "In Progress",
    "in_review": "In Review",
    "done": "Done",
}

PRIORITY_OPTIONS = ["lowest", "low", "medium", "high", "highest"]
PRIORITY_ICONS = {
    "lowest": "🔽",
    "low": "🔻",
    "medium": "➖",
    "high": "🔺",
    "highest": "🔴",
}

# Issue types selectable when creating a top-level item.
# "subtask" is created only from within a parent task, not from this list.
ISSUE_TYPE_ROOT_OPTIONS = ["task", "story", "bug"]
ISSUE_TYPE_ICONS = {
    "task": "✅",
    "subtask": "↳",
    "bug": "🐞",
    "story": "📘",
}


# Overview categories derived from a task's status + due date (see task_category).
# Colour None = default text colour.
TASK_CATEGORY_COLORS = {
    "in_progress": None,
    "completed": "green",
    "overdue": "red",
}


def task_category(task: dict, today: dt.date = None) -> str:
    """Completed if done; Overdue if the due date has passed and it isn't done;
    otherwise In Progress (covers To Do / In Progress / In Review)."""
    if task["status"] == "done":
        return "completed"
    today = today or dt.date.today()
    due = task.get("due_date")
    if due and due < today.isoformat():
        return "overdue"
    return "in_progress"


def task_date_label(task: dict) -> str:
    """'Completed: <date>' for done tasks, else 'Due: <date>'."""
    if task["status"] == "done":
        done_on = (task.get("completed_at") or "")[:10]
        return f"Completed: {done_on or '—'}"
    return f"Due: {task.get('due_date') or '—'}"


# ---------------------------------------------------------------
# Digital assets payment tracker
# ---------------------------------------------------------------
SERVICE_TYPE_OPTIONS = ["server", "domain", "email", "vendor_payment", "others"]
SERVICE_TYPE_LABELS = {
    "server": "Server",
    "domain": "Domain",
    "email": "Email",
    "vendor_payment": "Vendor payment",
    "others": "Others",
}

BILLING_CYCLE_OPTIONS = ["monthly", "quarterly", "semi_annually", "annually"]
BILLING_CYCLE_LABELS = {
    "monthly": "Monthly",
    "quarterly": "Quarterly",
    "semi_annually": "Semi-annually",
    "annually": "Annually",
}
BILLING_CYCLE_MONTHS = {"monthly": 1, "quarterly": 3, "semi_annually": 6, "annually": 12}

# Invoices due within this many days trigger a portal/email reminder.
ASSET_REMINDER_DAYS = 7
ASSET_STATUS_LABELS = {
    "upcoming": "Upcoming",
    "due_soon": "Due this week",
    "overdue": "Overdue",
}
ASSET_STATUS_COLORS = {
    "upcoming": None,
    "due_soon": "orange",
    "overdue": "red",
}


def service_type_label(asset: dict) -> str:
    """The typed-in text for 'others', else the dropdown label."""
    if asset["service_type"] == "others" and asset.get("service_type_other"):
        return asset["service_type_other"]
    return SERVICE_TYPE_LABELS.get(asset["service_type"], asset["service_type"])


def asset_label(asset: dict) -> str:
    return f"{asset['entity']} — {service_type_label(asset)}"


def due_when(days: int) -> str:
    """Human phrase for days until due: 'overdue by 3 day(s)', 'due today', 'due in 5 day(s)'."""
    if days < 0:
        return f"overdue by {-days} day(s)"
    if days == 0:
        return "due today"
    return f"due in {days} day(s)"


def asset_days_left(asset: dict, today: dt.date = None) -> int:
    today = today or dt.date.today()
    return (dt.date.fromisoformat(asset["due_date"]) - today).days


def asset_status(asset: dict, today: dt.date = None) -> str:
    """Overdue if the due date has passed; Due this week if within
    ASSET_REMINDER_DAYS; otherwise Upcoming."""
    days = asset_days_left(asset, today)
    if days < 0:
        return "overdue"
    if days <= ASSET_REMINDER_DAYS:
        return "due_soon"
    return "upcoming"


def add_months(d: dt.date, months: int) -> dt.date:
    """Add calendar months, clamping the day to the target month's length
    (e.g. 31 Jan + 1 month -> 28/29 Feb)."""
    month_index = d.month - 1 + months
    year, month = d.year + month_index // 12, month_index % 12 + 1
    day = min(d.day, calendar.monthrange(year, month)[1])
    return dt.date(year, month, day)


def next_due_date(due: dt.date, billing_cycle: str) -> dt.date:
    """Advance a due date by one billing cycle."""
    return add_months(due, BILLING_CYCLE_MONTHS[billing_cycle])


# ---------------------------------------------------------------
# Compliance tracker
# ---------------------------------------------------------------
ENTITY_KIND_OPTIONS = ["company", "individual"]
ENTITY_KIND_LABELS = {"company": "Company", "individual": "Individual"}

# Company register fields. Dropdowns also offer "Other…" for anything new.
COMPANY_COUNTRY_OPTIONS = ["India", "Dubai", "Canada"]
COMPANY_TYPE_OPTIONS = ["Pvt Ltd", "LLP", "Partnership firm", "Trust", "LLC"]
CURRENCY_OPTIONS = ["AED", "USD", "INR", "CAD", "EUR", "GBP"]
CURRENCY_BY_COUNTRY = {"India": "INR", "Dubai": "AED", "Canada": "CAD"}


def years_since(d, today: dt.date | None = None) -> str:
    """Age since a date as 'N yrs M mos' (recomputed every time, so it keeps
    up as time passes). '' for no date."""
    if not d:
        return ""
    if isinstance(d, str):
        d = dt.date.fromisoformat(d)
    today = today or dt.date.today()
    months = (today.year - d.year) * 12 + (today.month - d.month) - (today.day < d.day)
    if months < 0:
        return ""
    y, m = divmod(months, 12)
    return f"{y} yr{'s' if y != 1 else ''} {m} mo{'s' if m != 1 else ''}"


def fmt_number(n) -> str:
    """710000 -> '710,000'; 2.5 -> '2.5'; None -> ''."""
    if n is None or n == "":
        return ""
    n = float(n)
    return f"{n:,.0f}" if n.is_integer() else f"{n:,.2f}".rstrip("0").rstrip(".")


def shareholder_table(shareholders) -> list:
    """Shareholders with calculated paid-up capital (shares x face value) and
    % of shareholding (shares / total shares of the listed shareholders)."""
    rows = [s for s in (shareholders or []) if s.get("name")]
    total = sum(float(s.get("shares") or 0) for s in rows)
    out = []
    for s in rows:
        shares, face = s.get("shares"), s.get("face_value")
        out.append({
            "name": s["name"],
            "shares": shares,
            "face_value": face,
            "paid_up": float(shares) * float(face) if shares is not None and face is not None else None,
            "pct": float(shares) / total * 100 if shares is not None and total else None,
        })
    return out


def register_lines(entity: dict) -> list:
    """Pair an entity's directors with its shareholders, one register row
    each: as many rows as the longer list (at least one). Returns
    [(director name or None, shareholder_table() row or None), ...]."""
    dirs = [d["name"] for d in (entity.get("directors") or []) if d.get("name")]
    shs = shareholder_table(entity.get("shareholders"))
    n = max(len(dirs), len(shs), 1)
    return [(dirs[i] if i < len(dirs) else None, shs[i] if i < len(shs) else None) for i in range(n)]


def entity_type_label(entity: dict) -> str:
    if entity.get("kind") == "company":
        return entity.get("entity_type") or "Company"
    return "Individual"


REGISTER_HEADERS = [
    "Sr. No.", "Name", "Country", "Type of entity", "Business Activity", "Incorporation Date",
    "Years", "Name of Director", "Name of Shareholder", "No. of Shares held", "Face Value",
    "Paid up share Capital/Contribution", "% of Shareholding", "Other details",
]
_REGISTER_WIDTHS = [7, 34, 12, 16, 28, 14, 14, 26, 26, 14, 14, 20, 12, 40]


def entity_register_xlsx(entities: list) -> bytes:
    """The entity register as an Excel file: one row per director /
    shareholder, real dates and numbers, and a Years formula that keeps
    counting from the incorporation date every time the file is opened."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    def num_fmt(v, currency=None):
        base = "#,##0" if float(v).is_integer() else "#,##0.00"
        return f'"{currency} "{base}' if currency else base

    wb = Workbook()
    ws = wb.active
    ws.title = "Entities"
    ws.append(REGISTER_HEADERS)
    col = {h: i + 1 for i, h in enumerate(REGISTER_HEADERS)}

    r = 1
    for n, e in enumerate(entities, start=1):
        cur = e.get("currency")
        for j, (director, s) in enumerate(register_lines(e)):
            r += 1
            if j == 0:
                ws.cell(r, col["Sr. No."], n)
                ws.cell(r, col["Name"], e["name"])
                ws.cell(r, col["Country"], e.get("country"))
                ws.cell(r, col["Type of entity"], entity_type_label(e))
                ws.cell(r, col["Business Activity"], e.get("business_activity"))
                if e.get("incorporation_date"):
                    c = ws.cell(r, col["Incorporation Date"], dt.date.fromisoformat(e["incorporation_date"]))
                    c.number_format = "DD.MM.YYYY"
                    f = f"F{r}"
                    ws.cell(r, col["Years"],
                            f'=DATEDIF({f},TODAY(),"y")&" yrs "&DATEDIF({f},TODAY(),"ym")&" mos"')
                ws.cell(r, col["Other details"], pairs_to_text(e.get("details")) or None)
            ws.cell(r, col["Name of Director"], director)
            if s:
                ws.cell(r, col["Name of Shareholder"], s["name"])
                for header, key, currency in (
                    ("No. of Shares held", "shares", None),
                    ("Face Value", "face_value", cur),
                    ("Paid up share Capital/Contribution", "paid_up", cur),
                ):
                    if s[key] is not None:
                        c = ws.cell(r, col[header], s[key])
                        c.number_format = num_fmt(s[key], currency)
                if s["pct"] is not None:
                    c = ws.cell(r, col["% of Shareholding"], s["pct"] / 100)
                    c.number_format = "0%" if float(s["pct"]).is_integer() else "0.00%"

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=1, max_row=r, max_col=len(REGISTER_HEADERS)):
        for c in row:
            c.border = border
            c.alignment = Alignment(vertical="top", wrap_text=True)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="FCE4D6")
    for i, w in enumerate(_REGISTER_WIDTHS, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

COMPLIANCE_STATUS_OPTIONS = ["pending", "in_progress", "done", "not_applicable"]
COMPLIANCE_STATUS_LABELS = {
    "pending": "Pending",
    "in_progress": "In progress",
    "done": "Done",
    "not_applicable": "N/A",
}
# Items in these statuses never raise reminders.
COMPLIANCE_CLOSED_STATUSES = {"done", "not_applicable"}

# None = not set yet (Varsha fills it in). Months to roll forward on "Mark done".
COMPLIANCE_FREQUENCY_OPTIONS = [
    None, "one_time", "monthly", "quarterly", "half_yearly", "annually",
    "every_2_years", "every_3_years", "every_5_years", "every_10_years",
]
COMPLIANCE_FREQUENCY_LABELS = {
    None: "Not set",
    "one_time": "One-time",
    "monthly": "Monthly",
    "quarterly": "Quarterly",
    "half_yearly": "Half-yearly",
    "annually": "Annually",
    "every_2_years": "Every 2 years",
    "every_3_years": "Every 3 years",
    "every_5_years": "Every 5 years",
    "every_10_years": "Every 10 years",
}
COMPLIANCE_FREQUENCY_MONTHS = {
    "monthly": 1, "quarterly": 3, "half_yearly": 6, "annually": 12,
    "every_2_years": 24, "every_3_years": 36, "every_5_years": 60, "every_10_years": 120,
}
COMPLIANCE_DEFAULT_REMINDER_DAYS = 30

# Alert derived from status + due date + the item's own reminder window.
COMPLIANCE_ALERT_COLORS = {
    "overdue": "red",
    "due_soon": "orange",
    "done": "green",
    "ok": None,
}


def compliance_days_left(item: dict, today: dt.date = None):
    """Days until due, or None if the item has no due date."""
    if not item.get("due_date"):
        return None
    today = today or dt.date.today()
    return (dt.date.fromisoformat(item["due_date"]) - today).days


def compliance_alert(item: dict, today: dt.date = None) -> str:
    """'done' for closed items; 'overdue' past the due date; 'due_soon' inside
    the item's reminder window; otherwise 'ok' (incl. no due date set)."""
    if item["status"] in COMPLIANCE_CLOSED_STATUSES:
        return "done"
    days = compliance_days_left(item, today)
    if days is None:
        return "ok"
    if days < 0:
        return "overdue"
    if days <= int(item.get("reminder_days") or 0):
        return "due_soon"
    return "ok"


def compliance_alert_label(item: dict, today: dt.date = None) -> str:
    alert = compliance_alert(item, today)
    if alert == "overdue" or alert == "due_soon":
        return due_when(compliance_days_left(item, today)).capitalize()
    return ""


def pairs_to_text(pairs) -> str:
    """[{'field': 'Policy no.', 'value': 'X1'}] -> 'Policy no.: X1'."""
    return "; ".join(
        f"{p.get('field')}: {p.get('value')}" for p in (pairs or [])
        if p.get("field") and str(p.get("value") or "").strip()
    )


def name_to_id_map(profiles: list) -> dict:
    return {p["full_name"]: p["id"] for p in profiles}


def checkbox_filter(st, label: str, options: list, key_prefix: str, format_func=None, per_row: int = 4) -> list:
    """Render `options` as a row of plain checkboxes (all checked by default)
    instead of a multiselect's colored tag pills. Returns the checked options."""
    st.caption(label)
    selected = []
    cols = st.columns(min(per_row, len(options)) or 1)
    for i, opt in enumerate(options):
        text = format_func(opt) if format_func else str(opt)
        if cols[i % len(cols)].checkbox(text, value=True, key=f"{key_prefix}_{opt}"):
            selected.append(opt)
    return selected


def project_progress(tasks: list, rollup: bool = False) -> dict:
    """Aggregate progress for one project's tasks (flat list, subtasks included).

    Hours completion = sum(logged_hours) / sum(estimate_hours) — reflects
    actual hours worked against what was estimated, independent of status.

    rollup=True treats a sub-task as a further breakdown of its parent's
    estimate rather than additional work on top of it: each root task
    contributes max(its own hours, sum of its sub-tasks' hours) instead of
    (its own hours + sum of its sub-tasks' hours). Use this for project-wide
    totals; leave it off (default) when `tasks` is an arbitrary subset (e.g.
    one person's tasks) where parent/child pairs may be split across people.
    """
    total_tasks = len(tasks)
    done_tasks = [t for t in tasks if t["status"] == "done"]
    task_pct = (len(done_tasks) / total_tasks * 100) if total_tasks else 0.0

    if rollup:
        children_by_parent: dict = {}
        for t in tasks:
            pid = t.get("parent_id")
            if pid:
                children_by_parent.setdefault(pid, []).append(t)

        total_hours = 0.0
        logged_hours = 0.0
        for t in tasks:
            if t.get("parent_id"):
                continue  # counted via its parent below
            children = children_by_parent.get(t["id"], [])
            child_est = sum(float(c.get("estimate_hours") or 0) for c in children)
            child_logged = sum(float(c.get("logged_hours") or 0) for c in children)
            total_hours += max(float(t.get("estimate_hours") or 0), child_est)
            logged_hours += max(float(t.get("logged_hours") or 0), child_logged)
    else:
        total_hours = sum(float(t.get("estimate_hours") or 0) for t in tasks)
        logged_hours = sum(float(t.get("logged_hours") or 0) for t in tasks)

    return {
        "total_tasks": total_tasks,
        "done_tasks": len(done_tasks),
        "task_pct": task_pct,
        "total_hours": total_hours,
        "logged_hours": logged_hours,
        "hours_pct": (logged_hours / total_hours * 100) if total_hours else 0.0,
    }


def per_assignee_progress(tasks: list, id_to_name: dict) -> list:
    """Group a project's tasks by assignee and compute progress per person.
    Returns a list of dicts sorted by full name, "Unassigned" last."""
    by_assignee: dict = {}
    for t in tasks:
        aid = t.get("assignee_id")
        by_assignee.setdefault(aid, []).append(t)

    rows = []
    for aid, ts in by_assignee.items():
        prog = project_progress(ts)
        rows.append({
            "name": id_to_name.get(aid, "Unassigned") if aid else "Unassigned",
            **prog,
        })

    rows.sort(key=lambda r: (r["name"] == "Unassigned", r["name"]))
    return rows
